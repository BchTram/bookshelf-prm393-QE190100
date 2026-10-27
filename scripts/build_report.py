"""Assemble the exam Word template and only existing, real screenshots.
Run with python-docx from the document runtime. Private metadata stays local.
The result is a draft until public URLs, all required screenshots and visual QA pass.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
CAPTIONS = [
    ('01_firebase_auth', 'Firebase console — Authentication'),
    ('02_firestore_data', 'Firebase console — Firestore data'),
    ('03_published_rules', 'Firebase console — Security Rules'),
    ('04_register_validation', 'Register screen — validation'),
    ('05_register_success', 'Register screen — success'),
    ('06_login_empty', 'Login screen'),
    ('07_list_after_login', 'List screen — after login'),
    ('08_add_validation', 'Add Book form — validation'),
    ('09_add_success', 'Add Book form — success'),
    ('10_filter_genre', 'Filters and search — genre filter'),
    ('11_filter_status', 'Filters and search — status filter'),
    ('12_filter_favorites', 'Filters and search — favorites-only switch'),
    ('13_search', 'Filters and search — search box'),
    ('14_summary', 'Summary'),
    ('15_detail_view', 'Detail screen — view mode'),
    ('16_detail_edit_validation', 'Detail screen — edit mode, validation'),
    ('17_detail_after_save', 'Detail screen — after Save'),
    ('18_favorite_before', 'Favorite toggle and Delete — before favorite toggle'),
    ('19_favorite_after', 'Favorite toggle and Delete — after favorite toggle'),
    ('20_delete_confirmation', 'Favorite toggle and Delete — before Delete'),
    ('21_delete_after', 'Favorite toggle and Delete — after Delete'),
    ('22_empty_fresh_account', 'Empty state'),
    ('00_home', 'Home screen — public features'),
]


def main():
    metadata_path = ROOT / '.private/submission_metadata.json'
    metadata = json.loads(metadata_path.read_text(encoding='utf-8-sig')) if metadata_path.exists() else {}
    credentials_path = ROOT / '.private/grader_account.json'
    credentials = json.loads(credentials_path.read_text()) if credentials_path.exists() else {}
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(8)
    entries = [
        ('Student Code', 'QE190100'), ('Full Name', 'Pham Ngo Bich Tram'),
        ('GitHub Repository URL', metadata.get('repository_url', 'https://REPLACE_WITH_PUBLIC_GITHUB_REPOSITORY')),
        ('APK Download Link', metadata.get('apk_url', 'https://REPLACE_WITH_PUBLIC_APK_DOWNLOAD')),
        ('Firebase Project ID', 'bookshelf-prm393-7135b'),
        ('Test account email', credentials.get('GRADER_EMAIL', '<real-grader-email>')),
        ('Test account password', credentials.get('GRADER_PASSWORD', '<real-grader-password>')),
    ]
    for label, value in entries:
        doc.add_paragraph(f'{label}: {value}')
    missing = []
    included = []
    for filename, caption in CAPTIONS:
        path = ROOT / 'report_screenshots' / f'{filename}.png'
        if not path.exists():
            missing.append(filename)
            continue
        doc.add_page_break()
        heading = doc.add_paragraph(caption)
        heading.paragraph_format.keep_with_next = True
        heading.runs[0].bold = True
        width, height = Image.open(path).size
        scale = min(6.8 / width, 8.8 / height)
        picture = doc.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(str(path), width=Inches(width * scale), height=Inches(height * scale))
        included.append(filename)
    folder = ROOT / 'submission'
    folder.mkdir(exist_ok=True)
    output = folder / 'QE190100.docx'
    doc.save(output)
    status = {'report': str(output), 'included_screenshots': included,
              'missing_screenshots': missing,
              'public_urls_configured': bool(metadata.get('repository_url') and metadata.get('apk_url')),
              'render_verified': False}
    (ROOT / '.private/report_status.json').write_text(json.dumps(status, indent=2))
    print(f'Word draft created; {len(included)} actual screenshots included; {len(missing)} still missing. Password is not printed.')


if __name__ == '__main__':
    main()
